from docx import Document
from sqlalchemy.orm import Session
from app.db.session import SessionLocal
from app.models.career_path import CareerPath
from app.models.skill import Skill
from app.models.question import Question

def create_word_report(db: Session) -> None:
    document = Document()
    document.add_heading('Skills and Questions by Career Path', 0)

    career_paths = db.query(CareerPath).all()

    for cp in career_paths:
        document.add_heading(cp.name, level=1)
        skills = db.query(Skill).filter(Skill.career_path_id == cp.id).all()
        for skill in skills:
            document.add_heading(skill.name, level=2)
            questions = db.query(Question).filter(Question.skill_id == skill.id).all()
            for question in questions:
                document.add_paragraph(f"• {question.text}")

    document.save('skills_and_questions_report.docx')
    print("✅ Word document generated: skills_and_questions_report.docx")

def main() -> None:
    db = SessionLocal()
    try:
        create_word_report(db)
    finally:
        db.close()

if __name__ == "__main__":
    main()